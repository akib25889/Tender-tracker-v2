import json
import os
import sys
from pathlib import Path
import openpyxl

# Add backend to path
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

from app.core.database import SessionLocal, Base, engine, run_migrations
from app.models.organization import Organization

# Path to the Excel directory
EXCEL_PATH = (
    backend_dir.parent
    / "project design"
    / "ministry"
    / "Bangladesh_Govt_Tender_Directory.xlsx"
)

# Ministry definitions and mapping to clean IDs
MINISTRY_METADATA = {
    "President's Office": {
        "id": "ORG-BD-BANGABHABAN",
        "short_name": "Bangabhaban",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Prime Minister's Office": {
        "id": "ORG-BD-PMO",
        "short_name": "PMO",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Armed Forces Division": {
        "id": "ORG-BD-AFD",
        "short_name": "AFD",
        "type": "DIVISION",
        "priority": "HIGH",
    },
    "Cabinet Division": {
        "id": "ORG-BD-CABINET",
        "short_name": "Cabinet Division",
        "type": "DIVISION",
        "priority": "CRITICAL",
    },
    "Ministry of Chittagong Hill Tracts Affairs": {
        "id": "ORG-BD-MOCHTA",
        "short_name": "MoCHTA",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Primary and Mass Education": {
        "id": "ORG-BD-MOPME",
        "short_name": "MoPME",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Agriculture": {
        "id": "ORG-BD-MOA",
        "short_name": "MoA",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Civil Aviation and Tourism": {
        "id": "ORG-BD-MOCAT",
        "short_name": "MoCAT",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Commerce": {
        "id": "ORG-BD-MINCOM",
        "short_name": "MinCom",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Road Transport and Bridges": {
        "id": "ORG-BD-MORTH",
        "short_name": "MoRTH",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Cultural Affairs": {
        "id": "ORG-BD-MOCA",
        "short_name": "MoCA",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Defence": {
        "id": "ORG-BD-MOD",
        "short_name": "MoD",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Food": {
        "id": "ORG-BD-MOFOOD",
        "short_name": "MoFood",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Education": {
        "id": "ORG-BD-MOEDU",
        "short_name": "MoEdu",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Power, Energy and Mineral Resources": {
        "id": "ORG-BD-MPEMR",
        "short_name": "MPEMR",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Environment, Forest and Climate Change": {
        "id": "ORG-BD-MOEF",
        "short_name": "MoEFCC",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Public Administration": {
        "id": "ORG-BD-MOPA",
        "short_name": "MoPA",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Fisheries and Livestock": {
        "id": "ORG-BD-MOFL",
        "short_name": "MoFL",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Finance": {
        "id": "ORG-BD-MOF",
        "short_name": "MoF",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Foreign Affairs": {
        "id": "ORG-BD-MOFA",
        "short_name": "MoFA",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Health and Family Welfare": {
        "id": "ORG-BD-MOHFW",
        "short_name": "MoHFW",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Home Affairs": {
        "id": "ORG-BD-MOHA",
        "short_name": "MoHA",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Housing and Public Works": {
        "id": "ORG-BD-MOHPW",
        "short_name": "MoHPW",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Industries": {
        "id": "ORG-BD-MOIND",
        "short_name": "MoInd",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Information and Broadcasting": {
        "id": "ORG-BD-MOI",
        "short_name": "MoI",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Textiles and Jute": {
        "id": "ORG-BD-MOTJ",
        "short_name": "MoTJ",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Labour and Employment": {
        "id": "ORG-BD-MOLE",
        "short_name": "MoLE",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Law, Justice and Parliamentary Affairs": {
        "id": "ORG-BD-MINLAW",
        "short_name": "MinLaw",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Land": {
        "id": "ORG-BD-MINLAND",
        "short_name": "MinLand",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Local Government, Rural Development and Co-operatives": {
        "id": "ORG-BD-LGRD",
        "short_name": "LGRD",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Planning": {
        "id": "ORG-BD-PLANCOMM",
        "short_name": "MoPlan",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Posts, Telecommunications and Information Technology": {
        "id": "ORG-BD-MOPTIT",
        "short_name": "MoPTIT",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Religious Affairs": {
        "id": "ORG-BD-MORA",
        "short_name": "MoRA",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Disaster Management and Relief": {
        "id": "ORG-BD-MODMR",
        "short_name": "MoDMR",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Shipping": {
        "id": "ORG-BD-MOS",
        "short_name": "MoS",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Social Welfare": {
        "id": "ORG-BD-MSW",
        "short_name": "MSW",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Women and Children Affairs": {
        "id": "ORG-BD-MOWCA",
        "short_name": "MoWCA",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Water Resources": {
        "id": "ORG-BD-MOWR",
        "short_name": "MoWR",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Youth and Sports": {
        "id": "ORG-BD-MOYSPORTS",
        "short_name": "MoYS",
        "type": "MINISTRY",
        "priority": "MEDIUM",
    },
    "Ministry of Liberation War Affairs": {
        "id": "ORG-BD-MOLWA",
        "short_name": "MoLWA",
        "type": "MINISTRY",
        "priority": "LOW",
    },
    "Ministry of Expatriates' Welfare and Overseas Employment": {
        "id": "ORG-BD-PROBASHI",
        "short_name": "MEWOE",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
    "Ministry of Railways": {
        "id": "ORG-BD-MOR",
        "short_name": "MoR",
        "type": "MINISTRY",
        "priority": "CRITICAL",
    },
    "Ministry of Science and Technology": {
        "id": "ORG-BD-MOST",
        "short_name": "MoST",
        "type": "MINISTRY",
        "priority": "HIGH",
    },
}

# Subordinate offices mapping: identifies exact parent ministry/division and clean ID & type
OFFICE_HIERARCHY_RULES = {
    "BBS": {
        "id": "ORG-BD-BBS",
        "parent_id": "ORG-BD-PLANCOMM",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "BBS",
    },
    "Chemical Corporation": {
        "id": "ORG-BD-BCIC",
        "parent_id": "ORG-BD-MOIND",
        "type": "CORPORATION",
        "priority": "HIGH",
        "short_name": "BCIC",
    },
    "BCSAA": {
        "id": "ORG-BD-BCSAA",
        "parent_id": "ORG-BD-MOPA",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "BCSAA",
    },
    "Land Port": {
        "id": "ORG-BD-BLPA",
        "parent_id": "ORG-BD-MOS",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "BLPA",
    },
    "Meteorological": {
        "id": "ORG-BD-BMD",
        "parent_id": "ORG-BD-MOD",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "BMD",
    },
    "Parjatan": {
        "id": "ORG-BD-BPC-CORP",
        "parent_id": "ORG-BD-MOCAT",
        "type": "CORPORATION",
        "priority": "MEDIUM",
        "short_name": "Parjatan Corp",
    },
    "PDB": {
        "id": "ORG-BD-BPDB",
        "parent_id": "ORG-BD-MPEMR",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "BPDB",
    },
    "PSC": {
        "id": "ORG-BD-BPSC",
        "parent_id": "ORG-BD-GOV",
        "type": "CONSTITUTIONAL_BODY",
        "priority": "HIGH",
        "short_name": "BPSC",
    },
    "SPARRSO": {
        "id": "ORG-BD-SPARRSO",
        "parent_id": "ORG-BD-MOD",
        "type": "AUTHORITY",
        "priority": "MEDIUM",
        "short_name": "SPARRSO",
    },
    "BTRC": {
        "id": "ORG-BD-BTRC",
        "parent_id": "ORG-BD-MOPTIT",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "BTRC",
    },
    "Water Development": {
        "id": "ORG-BD-BWDB",
        "parent_id": "ORG-BD-MOWR",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "BWDB",
    },
    "Archive": {
        "id": "ORG-BD-NANL",
        "parent_id": "ORG-BD-MOCA",
        "type": "DEPARTMENT",
        "priority": "LOW",
        "short_name": "NANL",
    },
    "Environment": {
        "id": "ORG-BD-DOE",
        "parent_id": "ORG-BD-MOEF",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DoE",
    },
    "Passport": {
        "id": "ORG-BD-DIP",
        "parent_id": "ORG-BD-MOHA",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DIP",
    },
    "Shipping": {
        "id": "ORG-BD-DOS",
        "parent_id": "ORG-BD-MOS",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DoS",
    },
    "DSS": {
        "id": "ORG-BD-DSS",
        "parent_id": "ORG-BD-MSW",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "DSS",
    },
    "Family Planning": {
        "id": "ORG-BD-DGFP",
        "parent_id": "ORG-BD-MOHFW",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "DGFP",
    },
    "DG Food": {
        "id": "ORG-BD-DGFOOD",
        "parent_id": "ORG-BD-MOFOOD",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "DG Food",
    },
    "DGHS": {
        "id": "ORG-BD-DGHS",
        "parent_id": "ORG-BD-MOHFW",
        "type": "DIRECTORATE",
        "priority": "CRITICAL",
        "short_name": "DGHS",
    },
    "DSHE": {
        "id": "ORG-BD-DSHE",
        "parent_id": "ORG-BD-MOEDU",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "DSHE",
    },
    "Forest": {
        "id": "ORG-BD-BFOREST",
        "parent_id": "ORG-BD-MOEF",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "Forest Dept",
    },
    "GTCL": {
        "id": "ORG-BD-GTCL",
        "parent_id": "ORG-BD-PETROBANGLA",
        "type": "STATE_OWNED_ENTERPRISE",
        "priority": "HIGH",
        "short_name": "GTCL",
    },
    "Land Survey": {
        "id": "ORG-BD-DLRS",
        "parent_id": "ORG-BD-MINLAND",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DLRS",
    },
    "Land Reform": {
        "id": "ORG-BD-LRB",
        "parent_id": "ORG-BD-MINLAND",
        "type": "AUTHORITY",
        "priority": "MEDIUM",
        "short_name": "LRB",
    },
    "LGED": {
        "id": "ORG-BD-LGED",
        "parent_id": "ORG-BD-LGRD",
        "type": "DEPARTMENT",
        "priority": "CRITICAL",
        "short_name": "LGED",
    },
    "NGO Bureau": {
        "id": "ORG-BD-NGOAB",
        "parent_id": "ORG-BD-PMO",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "NGOAB",
    },
    "PKSF": {
        "id": "ORG-BD-PKSF",
        "parent_id": "ORG-BD-MOF",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "PKSF",
    },
    "Petrobangla": {
        "id": "ORG-BD-PETROBANGLA",
        "parent_id": "ORG-BD-MPEMR",
        "type": "CORPORATION",
        "priority": "CRITICAL",
        "short_name": "Petrobangla",
    },
    "UGC": {
        "id": "ORG-BD-UGC",
        "parent_id": "ORG-BD-MOEDU",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "UGC",
    },
    "Planning": {
        "id": "ORG-BD-PLANDIV",
        "parent_id": "ORG-BD-PLANCOMM",
        "type": "DIVISION",
        "priority": "HIGH",
        "short_name": "Planning Div",
    },
    "EC": {
        "id": "ORG-BD-EC",
        "parent_id": "ORG-BD-GOV",
        "type": "CONSTITUTIONAL_BODY",
        "priority": "CRITICAL",
        "short_name": "EC",
    },
    "BPC": {
        "id": "ORG-BD-BPC-PETROL",
        "parent_id": "ORG-BD-MPEMR",
        "type": "CORPORATION",
        "priority": "CRITICAL",
        "short_name": "BPC (Petroleum)",
    },
    "DMTCL": {
        "id": "ORG-BD-DMTCL",
        "parent_id": "ORG-BD-MORTH",
        "type": "STATE_OWNED_ENTERPRISE",
        "priority": "CRITICAL",
        "short_name": "DMTCL",
    },
    "Roads and Highway": {
        "id": "ORG-BD-RHD",
        "parent_id": "ORG-BD-MORTH",
        "type": "DEPARTMENT",
        "priority": "CRITICAL",
        "short_name": "RHD",
    },
    "Disaster Deparment": {
        "id": "ORG-BD-DDM",
        "parent_id": "ORG-BD-MODMR",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DDM",
    },
    "DGME": {
        "id": "ORG-BD-DGME",
        "parent_id": "ORG-BD-MOHFW",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "DGME",
    },
    "Youth Development": {
        "id": "ORG-BD-DYD",
        "parent_id": "ORG-BD-MOYSPORTS",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "DYD",
    },
    "BANBEISE": {
        "id": "ORG-BD-BANBEIS",
        "parent_id": "ORG-BD-MOEDU",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "BANBEIS",
    },
    "Madrasa": {
        "id": "ORG-BD-DME",
        "parent_id": "ORG-BD-MOEDU",
        "type": "DIRECTORATE",
        "priority": "MEDIUM",
        "short_name": "DME",
    },
    "Hi-Tech Park": {
        "id": "ORG-BD-BHTPA",
        "parent_id": "ORG-BD-MOPTIT",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "BHTPA",
    },
    "ICT": {
        "id": "ORG-BD-DOICT",
        "parent_id": "ORG-BD-MOPTIT",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DoICT",
    },
    "Jatiyo Mahila Sangstha": {
        "id": "ORG-BD-JMS",
        "parent_id": "ORG-BD-MOWCA",
        "type": "AUTHORITY",
        "priority": "MEDIUM",
        "short_name": "JMS",
    },
    "Anti Corruption Commission": {
        "id": "ORG-BD-ACC",
        "parent_id": "ORG-BD-GOV",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "ACC",
    },
    "AG": {
        "id": "ORG-BD-OCAG",
        "parent_id": "ORG-BD-GOV",
        "type": "CONSTITUTIONAL_BODY",
        "priority": "CRITICAL",
        "short_name": "OCAG",
    },
    "DAE": {
        "id": "ORG-BD-DAE",
        "parent_id": "ORG-BD-MOA",
        "type": "DEPARTMENT",
        "priority": "HIGH",
        "short_name": "DAE",
    },
    "Food Safety Authority": {
        "id": "ORG-BD-BFSA",
        "parent_id": "ORG-BD-MOFOOD",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "BFSA",
    },
    "BADC": {
        "id": "ORG-BD-BADC",
        "parent_id": "ORG-BD-MOA",
        "type": "CORPORATION",
        "priority": "HIGH",
        "short_name": "BADC",
    },
    "DTCA": {
        "id": "ORG-BD-DTCA",
        "parent_id": "ORG-BD-MORTH",
        "type": "AUTHORITY",
        "priority": "HIGH",
        "short_name": "DTCA",
    },
    "Public Health Engineering": {
        "id": "ORG-BD-DPHE",
        "parent_id": "ORG-BD-LGRD",
        "type": "DEPARTMENT",
        "priority": "CRITICAL",
        "short_name": "DPHE",
    },
    "WASA": {
        "id": "ORG-BD-DWASA",
        "parent_id": "ORG-BD-LGRD",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "Dhaka WASA",
    },
    "Primary Education": {
        "id": "ORG-BD-DPE",
        "parent_id": "ORG-BD-MOPME",
        "type": "DIRECTORATE",
        "priority": "HIGH",
        "short_name": "DPE",
    },
    "UCEP": {
        "id": "ORG-BD-UCEP",
        "parent_id": "ORG-BD-MOEDU",
        "type": "OTHER",
        "priority": "MEDIUM",
        "short_name": "UCEP",
    },
    "IIFC": {
        "id": "ORG-BD-IIFC",
        "parent_id": "ORG-BD-MOF",
        "type": "STATE_OWNED_ENTERPRISE",
        "priority": "HIGH",
        "short_name": "IIFC",
    },
    "Explosives": {
        "id": "ORG-BD-EXPLOSIVES",
        "parent_id": "ORG-BD-MPEMR",
        "type": "DEPARTMENT",
        "priority": "MEDIUM",
        "short_name": "Explosives Dept",
    },
    "WARPO": {
        "id": "ORG-BD-WARPO",
        "parent_id": "ORG-BD-MOWR",
        "type": "AUTHORITY",
        "priority": "MEDIUM",
        "short_name": "WARPO",
    },
    "CEGIS": {
        "id": "ORG-BD-CEGIS",
        "parent_id": "ORG-BD-MOWR",
        "type": "OTHER",
        "priority": "HIGH",
        "short_name": "CEGIS",
    },
    "IWM": {
        "id": "ORG-BD-IWM",
        "parent_id": "ORG-BD-MOWR",
        "type": "OTHER",
        "priority": "HIGH",
        "short_name": "IWM",
    },
    "CAAB": {
        "id": "ORG-BD-CAAB",
        "parent_id": "ORG-BD-MOCAT",
        "type": "AUTHORITY",
        "priority": "CRITICAL",
        "short_name": "CAAB",
    },
}


def build_description(
    address_or_location: str, notice_url: str, tender_url: str
) -> str:
    parts = []
    if address_or_location:
        parts.append(f"Location: {address_or_location}")
    if notice_url:
        parts.append(f"Notice Portal: {notice_url}")
    if tender_url:
        parts.append(f"Tender Portal: {tender_url}")
    return " | ".join(parts)


def build_aliases(name: str, short_name: str, key: str = "") -> list:
    aliases = set()
    if short_name:
        aliases.add(short_name)
    if key and key != short_name:
        aliases.add(key)
    # Extract acronyms from parentheses e.g. "Local Government Engineering Department (LGED)"
    if "(" in name and ")" in name:
        inside = name[name.find("(") + 1 : name.find(")")].strip()
        if inside:
            aliases.add(inside)
    return [a for a in aliases if a]


SEED_JSON_PATH = backend_dir / "app" / "services" / "organizations_seed.json"


def generate_organization_records():
    # If Excel is absent (e.g., gitignored on production server), load from precompiled seed JSON
    if not EXCEL_PATH.exists():
        if SEED_JSON_PATH.exists():
            with open(SEED_JSON_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        raise FileNotFoundError(f"Directory file not found at: {EXCEL_PATH}")

    records = []

    # 1. Apex Root Organization
    records.append(
        {
            "id": "ORG-BD-GOV",
            "name": "Government of the People's Republic of Bangladesh",
            "short_name": "BD GOV",
            "type": "GOVERNMENT",
            "parent_id": None,
            "country": "Bangladesh",
            "website": "https://bangladesh.gov.bd",
            "priority": "CRITICAL",
            "aliases_json": json.dumps(
                ["BD GOV", "Bangladesh Government", "GoB", "National Portal"]
            ),
            "description": "Apex sovereign entity representing all procuring ministries, divisions, and directorates of Bangladesh.",
        }
    )

    wb = openpyxl.load_workbook(EXCEL_PATH)
    s_min = wb["Ministries"]
    s_off = wb["Govt Offices"]

    # 2. Ministries & Divisions (Parent: ORG-BD-GOV)
    for r in list(s_min.iter_rows(values_only=True))[3:]:
        if not r[1]:
            continue
        min_name = r[1].strip()
        address = r[2].strip() if r[2] else ""
        website = r[3].strip() if r[3] else ""
        notice_url = r[4].strip() if r[4] else ""
        tender_url = r[5].strip() if r[5] else ""

        meta = MINISTRY_METADATA.get(min_name)
        if not meta:
            # Fallback if name slightly differs
            clean_id = f"ORG-BD-{min_name.split()[-1].upper()}"
            meta = {
                "id": clean_id,
                "short_name": min_name.split()[-1],
                "type": "MINISTRY",
                "priority": "HIGH",
            }

        aliases = build_aliases(min_name, meta["short_name"])
        desc = build_description(address, notice_url, tender_url)

        records.append(
            {
                "id": meta["id"],
                "name": min_name,
                "short_name": meta["short_name"],
                "type": meta["type"],
                "parent_id": "ORG-BD-GOV",
                "country": "Bangladesh",
                "website": website or None,
                "priority": meta["priority"],
                "aliases_json": json.dumps(aliases) if aliases else None,
                "description": desc or None,
            }
        )

    # 3. Subordinate Offices, Directorates, Corporations & Boards
    for r in list(s_off.iter_rows(values_only=True))[3:]:
        if not r[1]:
            continue
        short_key = r[0].strip() if r[0] else ""
        off_name = r[1].strip()
        location = r[2].strip() if r[2] else ""
        website = r[3].strip() if r[3] else ""
        notice_url = r[4].strip() if r[4] else ""
        tender_url = r[5].strip() if r[5] else ""

        rule = OFFICE_HIERARCHY_RULES.get(short_key)
        if not rule:
            # Fallback match on substring of name
            for k, v in OFFICE_HIERARCHY_RULES.items():
                if k.lower() in off_name.lower():
                    rule = v
                    break

        if not rule:
            clean_id = f"ORG-BD-{short_key.replace(' ', '').upper()}"
            rule = {
                "id": clean_id,
                "parent_id": "ORG-BD-GOV",
                "type": "DEPARTMENT",
                "priority": "MEDIUM",
                "short_name": short_key,
            }

        aliases = build_aliases(off_name, rule["short_name"], short_key)
        desc = build_description(location, notice_url, tender_url)

        records.append(
            {
                "id": rule["id"],
                "name": off_name,
                "short_name": rule["short_name"],
                "type": rule["type"],
                "parent_id": rule["parent_id"],
                "country": "Bangladesh",
                "website": website or None,
                "priority": rule["priority"],
                "aliases_json": json.dumps(aliases) if aliases else None,
                "description": desc or None,
            }
        )

    return records


def import_ministry_hierarchy(db=None):
    Base.metadata.create_all(bind=engine)
    run_migrations()

    close_session = False
    if db is None:
        db = SessionLocal()
        close_session = True

    try:
        print("\n" + "=" * 65)
        print("  Bangladesh Procuring Organizations & Hierarchy Importer")
        print("=" * 65)

        records = generate_organization_records()
        print(f"Generated {len(records)} organization records from directory.")

        # Deduplicate records by ID
        unique_records = {}
        for r in records:
            unique_records[r["id"]] = r

        print(f"Unique organization entities to seed: {len(unique_records)}")

        # Validate parent references
        all_ids = set(unique_records.keys())
        for org_id, r in unique_records.items():
            if r["parent_id"] and r["parent_id"] not in all_ids:
                print(
                    f"WARNING: Parent {r['parent_id']} not found for {org_id} ({r['name']})! Falling back to ORG-BD-GOV."
                )
                r["parent_id"] = "ORG-BD-GOV"

        # Check existing and upsert
        added_count = 0
        updated_count = 0

        for r in unique_records.values():
            existing = db.query(Organization).filter(Organization.id == r["id"]).first()
            if existing:
                existing.name = r["name"]
                existing.short_name = r["short_name"]
                existing.type = r["type"]
                existing.parent_id = r["parent_id"]
                existing.country = r["country"]
                existing.website = r["website"]
                existing.priority = r["priority"]
                existing.aliases_json = r["aliases_json"]
                existing.description = r["description"]
                updated_count += 1
            else:
                db_org = Organization(
                    id=r["id"],
                    name=r["name"],
                    short_name=r["short_name"],
                    type=r["type"],
                    parent_id=r["parent_id"],
                    country=r["country"],
                    website=r["website"],
                    priority=r["priority"],
                    aliases_json=r["aliases_json"],
                    description=r["description"],
                )
                db.add(db_org)
                added_count += 1

        db.commit()
        print(
            f"Successfully committed: {added_count} created, {updated_count} updated."
        )
        print(f"Total organizations currently in DB: {db.query(Organization).count()}")
        print("=" * 65 + "\n")

    except Exception as e:
        db.rollback()
        print(f"ERROR during organization import: {e}")
        raise e
    finally:
        if close_session:
            db.close()


if __name__ == "__main__":
    import_ministry_hierarchy()
